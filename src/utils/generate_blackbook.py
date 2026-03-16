from __future__ import annotations

import argparse
from pathlib import Path

import joblib
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_TITLE = "AI-Driven Smart Healthcare Analytics: Diabetes Risk Prediction System"
COLLEGE_NAME = "Mulund College of Commerce (Autonomous)"
ACADEMIC_YEAR = "2025-26"
CODE_FILE_EXTENSIONS = {
    ".py",
    ".html",
    ".css",
    ".js",
    ".md",
    ".txt",
    ".yml",
    ".yaml",
    ".json",
    ".svg",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate SCT guideline-compliant black book.")
    parser.add_argument("--student-name", default="________________________", help="Student full name")
    parser.add_argument("--roll-number", default="________________________", help="Student roll number")
    parser.add_argument("--class-name", default="TYBSc/ TYBCA", help="Class name")
    parser.add_argument("--department", default="School of Computing & Technology", help="Department")
    parser.add_argument("--guide-name", default="________________________", help="Guide name")
    parser.add_argument("--guide-designation", default="Project Guide", help="Guide designation")
    parser.add_argument(
        "--output",
        default="docs/Smart_Healthcare_Analytics_Blackbook.docx",
        help="Output docx path",
    )
    return parser.parse_args()


def set_page_margins(section) -> None:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.left_margin = Inches(1.5)


def set_document_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(16 if level == 1 else 12)
        run.bold = True


def add_paragraph(
    doc: Document,
    text: str,
    *,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.JUSTIFY,
    bold: bool = False,
    italic: bool = False,
    font_size: int = 12,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
        for run in header_cells[index].paragraphs[0].runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            for run in cells[index].paragraphs[0].runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def set_page_number_start(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def configure_main_section_header_footer(section) -> None:
    header = section.header
    if not header.paragraphs:
        header.add_paragraph()
    header_para = header.paragraphs[0]
    header_para.clear()
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_para.add_run(PROJECT_TITLE)
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(12)
    header_run.italic = True

    footer = section.footer
    if not footer.paragraphs:
        footer.add_paragraph()
    footer_para = footer.paragraphs[0]
    footer_para.clear()
    footer_para.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
    footer_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_left = footer_para.add_run(f"{COLLEGE_NAME} ({ACADEMIC_YEAR})")
    footer_left.font.name = "Times New Roman"
    footer_left.font.size = Pt(12)
    footer_left.italic = True
    footer_para.add_run("\t")
    footer_right = footer_para.add_run("Page ")
    footer_right.font.name = "Times New Roman"
    footer_right.font.size = Pt(12)
    footer_right.italic = True
    add_page_number(footer_para)


def clean_route_line(line: str) -> str:
    line = line.strip().replace('"', "").replace("'", "")
    return line


def load_project_data(root: Path) -> dict:
    data_path = root / "data" / "diabetes.csv"
    artifact_path = root / "models" / "diabetes_model.joblib"
    app_path = root / "src" / "app.py"

    info: dict = {
        "dataset_rows": 0,
        "dataset_columns": 0,
        "outcome_counts": {},
        "means": {},
        "model": {},
        "routes": [],
        "images": [],
    }

    if data_path.exists():
        frame = pd.read_csv(data_path)
        info["dataset_rows"] = int(len(frame))
        info["dataset_columns"] = int(len(frame.columns))
        if "Outcome" in frame.columns:
            info["outcome_counts"] = {
                str(int(k)): int(v) for k, v in frame["Outcome"].value_counts().to_dict().items()
            }
        info["means"] = {
            key: float(value)
            for key, value in frame.mean(numeric_only=True).round(4).to_dict().items()
        }

    if artifact_path.exists():
        artifact = joblib.load(artifact_path)
        info["model"] = {
            "selected_model": artifact.get("selected_model", "unknown"),
            "metrics": artifact.get("metrics", {}),
            "cv_roc_auc": artifact.get("cv_roc_auc"),
            "feature_names": artifact.get("feature_names", []),
        }

    if app_path.exists():
        routes: list[str] = []
        for raw_line in app_path.read_text(encoding="utf-8").splitlines():
            stripped = raw_line.strip()
            if stripped.startswith("@app."):
                routes.append(clean_route_line(stripped))
        info["routes"] = routes

    image_dir = root / "docs" / "images"
    if image_dir.exists():
        info["images"] = sorted(str(path) for path in image_dir.glob("*.png"))

    return info


def add_code_line(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def safe_read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="latin-1")
        except UnicodeDecodeError:
            return path.read_text(encoding="utf-8", errors="ignore")


def collect_code_files(root: Path) -> list[Path]:
    files: list[Path] = []
    explicit_files = [
        root / "README.md",
        root / "requirements.txt",
        root / "environment.yml",
        root / "setup_environment.py",
    ]
    include_dirs = [root / "src", root / "tests"]
    skip_dir_names = {
        ".git",
        "venv",
        "__pycache__",
        ".ruff_cache",
        ".tmp",
    }

    for candidate in explicit_files:
        if candidate.exists() and candidate.is_file():
            files.append(candidate)

    for base in include_dirs:
        if not base.exists():
            continue
        for path in sorted(base.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in CODE_FILE_EXTENSIONS:
                continue
            if any(
                part in skip_dir_names or part.startswith("pytest-cache-files-")
                for part in path.parts
            ):
                continue
            files.append(path)

    unique_files: list[Path] = []
    seen: set[Path] = set()
    for path in files:
        resolved = path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique_files.append(path)
    return unique_files


def add_extended_website_documentation(doc: Document, project_data: dict) -> None:
    doc.add_page_break()
    add_heading(doc, "DETAILED WEBSITE DOCUMENTATION", level=1)
    add_heading(doc, "A. Website Purpose and User Segments", level=2)
    add_paragraph(
        doc,
        "The website is designed as a complete healthcare analytics dashboard for diabetes risk prediction. "
        "It supports three practical user intents: first, quick diabetes risk estimation through a guided form; "
        "second, learning-oriented analytics through interactive charts and model insights; and third, continuity "
        "of usage through account-based history and downloadable prediction files.",
    )
    add_paragraph(
        doc,
        "The interface is intentionally implemented in a dark theme to reduce visual fatigue during long usage, "
        "with high-contrast cards and action buttons so major workflows (login, predict, download, and review) "
        "remain obvious at first glance.",
    )

    add_heading(doc, "B. End-to-End User Journey", level=2)
    journey_steps = [
        "Landing and Hero: User opens the home route and sees project overview, core capabilities, and call-to-action buttons.",
        "Authentication: User can register a fresh account and then login with secure credential checks.",
        "Prediction Flow: User enters eight clinical parameters and submits the form for model inference.",
        "Result Feedback: System returns class label plus risk probability and stores it in prediction history.",
        "Export Flow: User can download latest prediction report (TXT/JSON/CSV) or entire history CSV.",
        "Analytics Flow: User explores charts (distribution, averages, feature-wise views) and patient records table.",
        "Model Insight Flow: User inspects selected model metrics and feature-importance perspective.",
    ]
    for step in journey_steps:
        add_paragraph(doc, f"- {step}", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "C. Detailed Navigation Map", level=2)
    add_table(
        doc,
        ["Navigation Item", "Destination", "Functional Role"],
        [
            ["Login", "/login", "Authenticates user and starts session"],
            ["Stats", "Dashboard statistics section", "Displays key aggregate metrics"],
            ["Charts", "Visual insights section", "Presents distribution and comparison charts"],
            ["Model", "Model insights section", "Shows algorithm performance and interpretation"],
            ["Patients", "Patient records section", "Lists data rows with filtering and pagination"],
            ["My Predictions", "History section", "Shows personal submissions and trend delta"],
            ["Predict Risk", "Prediction form section", "Primary interaction for diabetes risk inference"],
        ],
    )

    add_heading(doc, "D. Frontend Engineering Breakdown", level=2)
    add_paragraph(
        doc,
        "The frontend is built with server-rendered HTML templates and JavaScript-driven API integration. "
        "Core UX goals are readability, low latency interactions, and predictable state transitions. "
        "Forms include helper text and domain range hints to reduce invalid submissions. "
        "Cards, sections, and tables use a consistent visual language so a new user can navigate without training.",
    )
    add_table(
        doc,
        ["Frontend Area", "Implementation Notes"],
        [
            ["Layout", "Single-page style dashboard with anchored sections and scroll navigation"],
            ["Theme", "Dark palette with accent gradients aligned with project branding"],
            ["Form Validation", "Client-side numeric checks plus backend validation for safety"],
            ["Charts", "Chart.js rendering from API payloads with labeled legends and axes"],
            ["Tables", "Pagination controls and status badges for quick scanning"],
            ["Download Controls", "Buttons invoke export endpoints and trigger direct file save"],
        ],
    )

    add_heading(doc, "E. Backend Engineering Breakdown", level=2)
    add_paragraph(
        doc,
        "The backend exposes Flask routes for template rendering, authentication, prediction inference, "
        "history retrieval, report analysis, and exports. Each sensitive route checks authentication state. "
        "Prediction endpoints validate payload schema and numeric constraints before model inference.",
    )
    add_paragraph(
        doc,
        "Business logic and persistence are separated in a pragmatic way: model calls are isolated from "
        "database I/O, while API handlers transform request data into explicit response objects for the UI.",
    )

    add_heading(doc, "F. Security and Data Integrity Considerations", level=2)
    security_points = [
        "Password hashing is used for account credential storage (no plain-text password persistence).",
        "Session-based checks protect private endpoints such as prediction history and report storage.",
        "Input validation covers type conversion, required fields, and acceptable clinical ranges.",
        "Database interactions use controlled SQL statements, reducing injection risk in dynamic requests.",
        "Download exports are generated server-side from validated data and authenticated context.",
    ]
    for point in security_points:
        add_paragraph(doc, f"- {point}", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "G. Route Inventory from Source Code", level=2)
    route_rows = [[str(i + 1), route] for i, route in enumerate(project_data.get("routes", []))]
    if route_rows:
        add_table(doc, ["Sr. No.", "Decorator Definition"], route_rows)
    else:
        add_paragraph(
            doc,
            "No Flask decorators were discovered during route scan. Verify src/app.py availability.",
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )


def add_database_interaction_details(doc: Document, root: Path) -> None:
    db_path = root / "data" / "user_interactions.db"
    doc.add_page_break()
    add_heading(doc, "DATABASE INTERACTION SPECIFICATION", level=1)
    add_paragraph(
        doc,
        "The project persists all user interaction events into an SQLite database. "
        "This includes authentication entities, prediction records, and report analysis submissions.",
    )
    add_table(
        doc,
        ["Entity", "Purpose", "Operational Impact"],
        [
            ["users", "Stores account identity and role metadata", "Enables authenticated and personalized workflows"],
            ["predictions", "Stores every submitted prediction with probability", "Supports history view and exports"],
            ["reports", "Stores report-analysis submissions", "Maintains longitudinal interaction trail"],
        ],
    )
    add_paragraph(
        doc,
        f"Database file location in this project: {db_path}",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        italic=True,
    )
    add_paragraph(
        doc,
        "Database workflow detail: when a prediction request succeeds, the response object is first prepared "
        "for immediate UI rendering, then a normalized row is inserted into the predictions table with timestamp "
        "and all model input fields for future retrieval and export.",
    )


def add_dataset_annexure(doc: Document, root: Path) -> None:
    data_path = root / "data" / "diabetes.csv"
    doc.add_page_break()
    add_heading(doc, "ANNEXURE: COMPLETE DATASET RECORDS", level=1)
    if not data_path.exists():
        add_paragraph(doc, "Dataset file not found at data/diabetes.csv", align=WD_ALIGN_PARAGRAPH.LEFT)
        return

    frame = pd.read_csv(data_path)
    add_paragraph(
        doc,
        f"Dataset source file: {data_path}. Total rows: {len(frame)}. Total columns: {len(frame.columns)}.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        doc,
        "Below listing provides row-wise values for all records used by the project so the black book includes "
        "complete data context as requested.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    header_line = " | ".join(frame.columns.astype(str).tolist())
    add_code_line(doc, f"Columns: {header_line}")
    for idx, row in frame.iterrows():
        values = [str(row[col]) for col in frame.columns]
        add_code_line(doc, f"{idx + 1:04d}: " + " | ".join(values))


def add_source_code_annexure(doc: Document, root: Path) -> None:
    files = collect_code_files(root)
    total_lines = 0
    file_line_rows: list[list[str]] = []

    for path in files:
        line_count = len(safe_read_text(path).splitlines())
        total_lines += line_count
        file_line_rows.append([str(path.relative_to(root)), str(line_count)])

    doc.add_page_break()
    add_heading(doc, "ANNEXURE: COMPLETE SOURCE CODE LISTING", level=1)
    add_paragraph(
        doc,
        "This annexure includes full source code files used in the project implementation. "
        "Every line is listed with line numbers for traceability and viva reference.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        doc,
        f"Total files listed: {len(files)} | Total lines listed: {total_lines}",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        bold=True,
    )
    add_table(doc, ["File Path", "Line Count"], file_line_rows)

    for file_path in files:
        doc.add_page_break()
        relative = file_path.relative_to(root)
        add_heading(doc, f"File: {relative}", level=2)
        code_text = safe_read_text(file_path)
        lines = code_text.splitlines()
        add_paragraph(
            doc,
            f"Listing {len(lines)} lines from {relative}",
            align=WD_ALIGN_PARAGRAPH.LEFT,
            italic=True,
        )
        for line_no, line in enumerate(lines, start=1):
            safe_line = line if line else ""
            add_code_line(doc, f"{line_no:04d}: {safe_line}")


def add_project_operation_manual(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "PROJECT EXECUTION AND DEMONSTRATION MANUAL", level=1)
    add_heading(doc, "A. Environment Setup", level=2)
    setup_steps = [
        "1. Create or activate Python virtual environment.",
        "2. Install dependencies using requirements.txt.",
        "3. Train model artifact if not already available in models/ directory.",
        "4. Start Flask application from src/app.py.",
        "5. Open browser on http://127.0.0.1:5000/ and verify /health endpoint.",
    ]
    for step in setup_steps:
        add_paragraph(doc, step, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "B. Demonstration Script for Viva", level=2)
    demo_steps = [
        "Register a new user and show successful login.",
        "Submit a low-risk and high-risk sample in prediction form.",
        "Open prediction history and show latest-vs-previous delta.",
        "Download latest predicted file and full CSV history.",
        "Display charts, model insights, and patient records filters.",
        "Show database persistence by refreshing and reloading history.",
    ]
    for step in demo_steps:
        add_paragraph(doc, f"- {step}", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "C. Common Issues and Troubleshooting", level=2)
    add_table(
        doc,
        ["Issue", "Likely Reason", "Fix"],
        [
            ["Login fails", "Incorrect credential or missing user row", "Re-register user and verify DB path"],
            ["Prediction endpoint unauthorized", "No active session", "Login first and retry submission"],
            ["Charts not visible", "Missing API payload or JS error", "Check browser console and /api endpoints"],
            ["History export empty", "No prior predictions for user", "Submit one prediction and retry download"],
            ["App not starting", "Port conflict or env mismatch", "Free port 5000 and verify Python interpreter"],
        ],
    )


def add_expanded_blackbook_content(doc: Document, root: Path, project_data: dict) -> None:
    add_extended_website_documentation(doc, project_data)
    add_database_interaction_details(doc, root)
    add_project_operation_manual(doc)
    add_dataset_annexure(doc, root)
    add_source_code_annexure(doc, root)

def add_cover_pages(doc: Document, args: argparse.Namespace) -> None:
    add_paragraph(doc, "PARLE TILAK VIDYALAYA ASSOCIATION'S", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "MULUND COLLEGE OF COMMERCE (AUTONOMOUS)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "SCHOOL OF COMPUTING & TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "A PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14)
    add_paragraph(doc, "ON", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14)
    add_paragraph(doc, PROJECT_TITLE.upper(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=16)
    add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        doc,
        "Submitted in partial fulfillment of the requirements for the award of the degree of",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(
        doc,
        "BACHELOR OF SCIENCE (COMPUTER SCIENCE / INFORMATION TECHNOLOGY / COMPUTER APPLICATIONS)",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
    )
    add_paragraph(doc, "OR", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "BACHELOR OF COMPUTER APPLICATIONS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, f"By: {args.student_name}", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, f"Roll Number: {args.roll_number}", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Under the guidance of", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, args.guide_name, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, args.guide_designation, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "MULUND 400080, MAHARASHTRA", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, f"ACADEMIC YEAR {ACADEMIC_YEAR}", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

    doc.add_page_break()
    add_heading(doc, "COLLEGE CERTIFICATE", level=1)
    add_paragraph(
        doc,
        "Certificate will be issued by the college. Insert signed and stamped certificate page here.",
    )

    doc.add_page_break()
    add_heading(doc, "LIVE PROJECT CERTIFICATE (IF APPLICABLE)", level=1)
    add_paragraph(
        doc,
        "This project is an academic capstone project. If the department requires a live-project certificate, insert it here.",
    )

    doc.add_page_break()
    add_heading(doc, "PROFORMA FOR PROJECT APPROVAL", level=1)
    add_paragraph(
        doc,
        "Attach the original copy of the approved project proposal as provided in department proforma.",
    )

    doc.add_page_break()
    add_heading(doc, "ACKNOWLEDGEMENT", level=1)
    add_paragraph(
        doc,
        "I sincerely thank the In-charge Principal, the School of Computing & Technology, "
        "my project guide, faculty members, and all supporting staff for their guidance and support. "
        "I also thank my family and friends for their encouragement throughout the project duration.",
    )

    doc.add_page_break()
    add_heading(doc, "DECLARATION", level=1)
    add_table(
        doc,
        ["Field", "Details"],
        [
            ["Student Name", args.student_name],
            ["Roll Number", args.roll_number],
            ["Class", args.class_name],
            ["Department", args.department],
            ["Project Title", PROJECT_TITLE],
        ],
    )
    add_paragraph(
        doc,
        "I hereby declare that this project report is my original work. All sources and references used in this report "
        "have been acknowledged properly. This work has not been copied in part or full from any other source.",
    )
    add_paragraph(doc, "Student Signature: ________________________", bold=True)

    doc.add_page_break()
    add_heading(doc, "ABSTRACT", level=1)
    add_paragraph(
        doc,
        "This project presents a complete machine-learning based healthcare analytics system for diabetes risk prediction. "
        "The solution includes data preprocessing, model training and evaluation, web-based deployment, database-backed "
        "user interaction tracking, and visual analytics. The system accepts clinical input parameters such as Glucose, BMI, "
        "Blood Pressure, Age, and Diabetes Pedigree Function to estimate risk. The deployment includes authentication, "
        "prediction history, downloadable reports, and data-driven charts, making the platform suitable for educational "
        "research and capstone demonstration.",
    )


def add_index_page(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "INDEX (SEMESTER VI CAPSTONE FORMAT)", level=1)
    add_table(
        doc,
        ["Sr. No.", "Chapter Scheme", "Pages"],
        [
            ["1", "CHAPTER 1: INTRODUCTION", "1-4"],
            ["2", "CHAPTER 2: SURVEY OF TECHNOLOGIES (10 Papers)", "5-9"],
            ["3", "CHAPTER 3: REQUIREMENTS AND ANALYSIS", "10-14"],
            ["4", "CHAPTER 4: SYSTEM DESIGN", "15-20"],
            ["5", "CHAPTER 5: IMPLEMENTATION AND TESTING", "21-30"],
            ["6", "CHAPTER 6: RESULTS", "31-36"],
            ["7", "CHAPTER 7: CONCLUSION", "37-39"],
            ["8", "REFERENCES / BIBLIOGRAPHY", "40-41"],
            ["9", "ANNEXURE", "42-45"],
        ],
    )
    add_paragraph(
        doc,
        "Note: Final page numbers can vary after printing and formatting adjustments.",
        italic=True,
    )


def add_main_content(doc: Document, project_data: dict) -> None:
    doc.add_page_break()
    add_heading(doc, "CHAPTER 1: INTRODUCTION", level=1)
    add_heading(doc, "1.1 Background", level=2)
    add_paragraph(
        doc,
        "Diabetes is a major chronic condition requiring early and continuous risk screening. "
        "This capstone builds an AI-driven analytical platform that predicts diabetes risk from clinical inputs "
        "and provides a full-stack user experience including prediction history and downloadable reports.",
    )
    add_heading(doc, "1.2 Objectives", level=2)
    objectives = [
        "Build an end-to-end ML workflow from data preprocessing to deployment.",
        "Compare baseline models and select the best model based on ROC-AUC.",
        "Develop secure user interaction modules (login, registration, prediction history).",
        "Implement a database-backed system for reports and prediction records.",
        "Provide visual analytics and downloadable prediction outputs.",
    ]
    for item in objectives:
        add_paragraph(doc, f"- {item}", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_heading(doc, "1.3 Purpose, Scope, and Applicability", level=2)
    add_paragraph(doc, "Purpose: To provide a practical healthcare AI screening tool for diabetes risk prediction.")
    add_paragraph(
        doc,
        "Scope: Covers data ingestion, model training, web API development, dashboard UI, and persistent user interactions.",
    )
    add_paragraph(
        doc,
        "Applicability: Useful for academic demonstrations, basic clinical triage support, and learning healthcare analytics workflows.",
    )
    add_heading(doc, "1.4 Organization of Report", level=2)
    add_paragraph(
        doc,
        "This report is organized into survey of technologies, requirements analysis, system design, implementation/testing, "
        "results, conclusion, references, and annexure sections.",
    )

    doc.add_page_break()
    add_heading(doc, "CHAPTER 2: SURVEY OF TECHNOLOGIES (10 SELECTED PAPERS)", level=1)
    survey_rows = [
        ["1", "Cox (1958) - Logistic Regression", "Established logistic modeling for binary outcomes", "Used as baseline model"],
        ["2", "Breiman (2001) - Random Forests", "Ensemble learning with strong generalization", "Used as alternate candidate"],
        ["3", "Pima Indians Dataset (NIDDK)", "Clinical benchmark dataset for diabetes prediction", "Core dataset reference"],
        ["4", "Fawcett (2006) - ROC Analysis", "Threshold-independent classifier evaluation", "Used for model comparison"],
        ["5", "Bishop (Pattern Recognition)", "ML fundamentals and model validation", "Guided feature handling"],
        ["6", "Hastie et al. (ESL)", "Statistical learning framework", "Supported model-selection process"],
        ["7", "Scikit-learn User Guide", "Pipeline and cross-validation best practices", "Implemented in training flow"],
        ["8", "Flask Documentation", "Web backend route architecture", "Used for API and templates"],
        ["9", "SQLite Documentation", "Lightweight relational persistence", "Enabled user-interaction database"],
        ["10", "Chart.js Documentation", "Data visualization in web dashboard", "Used for analytics charts"],
    ]
    add_table(doc, ["Sr", "Paper / Source", "Key Contribution", "Project Relevance"], survey_rows)
    add_paragraph(
        doc,
        "The selected references collectively support model development, evaluation, deployment, and visualization strategy for the capstone system.",
    )

    doc.add_page_break()
    add_heading(doc, "CHAPTER 3: REQUIREMENTS AND ANALYSIS", level=1)
    add_heading(doc, "3.1 Problem Definition", level=2)
    add_paragraph(
        doc,
        "Manual diabetes risk assessment is time-consuming and not always accessible. "
        "The project solves this through a real-time predictive system with historical tracking and reporting.",
    )
    add_heading(doc, "3.2 Requirements Specification", level=2)
    add_table(
        doc,
        ["Type", "Requirement"],
        [
            ["Functional", "User registration and login"],
            ["Functional", "Predict diabetes risk from 8 clinical features"],
            ["Functional", "Save and retrieve prediction history per user"],
            ["Functional", "Upload/analyze report files and save output"],
            ["Functional", "Download individual prediction report and full CSV export"],
            ["Non-Functional", "Fast API response for inference requests"],
            ["Non-Functional", "Data persistence using SQLite database"],
            ["Non-Functional", "Readable dashboard with charts and tables"],
        ],
    )
    add_heading(doc, "3.3 Planning and Scheduling", level=2)
    add_table(
        doc,
        ["Phase", "Major Tasks", "Status"],
        [
            ["Phase 1", "Dataset preparation and preprocessing utilities", "Completed"],
            ["Phase 2", "Model training pipeline and selection", "Completed"],
            ["Phase 3", "Flask API and dashboard integration", "Completed"],
            ["Phase 4", "User interaction DB and report downloads", "Completed"],
            ["Phase 5", "Testing, documentation, and black book", "Completed"],
        ],
    )
    add_heading(doc, "3.4 Software and Hardware Requirements", level=2)
    add_table(
        doc,
        ["Category", "Details"],
        [
            ["Programming Language", "Python 3.13"],
            ["Backend", "Flask + flask-cors"],
            ["ML Libraries", "scikit-learn, pandas, numpy, joblib"],
            ["Frontend", "HTML/CSS/JavaScript, Chart.js"],
            ["Database", "SQLite"],
            ["OS", "Windows 10/11 (project tested on Windows)"],
            ["Minimum Hardware", "8 GB RAM, dual-core CPU, 500 MB free storage"],
        ],
    )
    add_heading(doc, "3.5 Preliminary Product Description", level=2)
    add_paragraph(
        doc,
        "The product is a web-based healthcare analytics dashboard with login/register, prediction engine, "
        "model insights, patient dataset exploration, report analysis, and downloadable prediction files.",
    )

    doc.add_page_break()
    add_heading(doc, "CHAPTER 4: SYSTEM DESIGN", level=1)
    add_heading(doc, "4.1 UML Diagrams", level=2)
    uml_items = [
        "Use Case Diagram: Patient login, predict risk, download reports.",
        "Activity Diagram: Input -> Validation -> Model Inference -> Save -> Download.",
        "Sequence Diagram: Browser -> Flask API -> Model -> SQLite -> Response.",
        "Class/Component View: Data utilities, training module, API layer, template layer.",
    ]
    for item in uml_items:
        add_paragraph(doc, f"- {item}", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_heading(doc, "4.2 Module Design", level=2)
    add_table(
        doc,
        ["Module", "Description"],
        [
            ["Data Factory", "Synthetic dataset generation fallback when real CSV is absent"],
            ["Preprocess", "Zero-to-missing normalization and split utilities"],
            ["Train Pipeline", "Model comparison, CV scoring, artifact generation"],
            ["Prediction API", "Input validation, risk probability, save interaction"],
            ["User Management", "Registration/login with hashed password storage"],
            ["Reports Module", "Upload analysis, report storage, retrieval"],
            ["Dashboard", "Stats, charts, patient table, prediction history"],
        ],
    )
    add_heading(doc, "4.3 Database Design", level=2)
    add_paragraph(
        doc,
        "The project uses SQLite for persistent user interaction tracking. Three core tables are designed:",
    )
    add_table(
        doc,
        ["Table", "Important Columns"],
        [
            ["users", "username (PK), password_hash, first_name, last_name, email, role, created_at"],
            ["predictions", "id (PK), username, submitted_at, prediction, risk_probability, patient fields, feature values"],
            ["reports", "id (PK), submitted_by, submitted_at, patient_name, age, clinical fields, notes"],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "CHAPTER 5: IMPLEMENTATION AND TESTING", level=1)
    add_heading(doc, "5.1 Overview of the Implementation", level=2)
    add_paragraph(
        doc,
        "Implementation follows modular software engineering principles. The backend is implemented in Flask, "
        "the ML pipeline in scikit-learn, and the frontend in vanilla web technologies.",
    )
    add_heading(doc, "5.2 Development Tools and Technology Stack", level=2)
    add_table(
        doc,
        ["Layer", "Tools / Files"],
        [
            ["Backend", "src/app.py (Flask routes, authentication, DB interaction)"],
            ["Model Training", "src/train.py (CV model selection, metrics, artifact save)"],
            ["Data Utilities", "src/utils/preprocess.py, src/utils/data_factory.py"],
            ["UI Templates", "src/templates/index.html, login.html, register.html"],
            ["Testing", "tests/test_train_and_app.py, tests/test_preprocess.py"],
            ["Database", "SQLite file: data/user_interactions.db"],
        ],
    )
    add_heading(doc, "5.3 System Architecture", level=2)
    add_paragraph(
        doc,
        "Client browser interacts with Flask endpoints. Flask validates input and invokes the trained model artifact. "
        "Predictions and reports are persisted in SQLite for future retrieval and export.",
    )
    add_heading(doc, "5.4 Module-wise Implementation", level=2)
    add_paragraph(doc, "5.4.1 User Interface Design: Dark themed responsive dashboard with role-neutral navigation.")
    add_paragraph(doc, "5.4.2 Functional Modules: Predict risk, analyze report, patient data exploration, charts.")
    add_paragraph(doc, "5.4.3 Database Integration: Users, predictions, and reports stored in SQLite.")
    add_paragraph(doc, "5.4.4 Input/Output Handling: JSON and multipart input with numeric and schema validation.")
    add_heading(doc, "5.5 Challenges Faced and Solutions Implemented", level=2)
    add_table(
        doc,
        ["Challenge", "Solution"],
        [
            ["Missing dataset in fresh environment", "Added deterministic synthetic dataset generator fallback"],
            ["Persistent user tracking requirement", "Migrated CSV interaction storage to SQLite database"],
            ["Test isolation with custom env paths", "Database path fallback linked to PREDICTIONS_PATH in tests"],
            ["Download requirement for predicted data", "Added single-file and full-history export endpoints"],
            ["UI consistency across pages", "Unified branding and color theme, shared logo asset"],
        ],
    )
    add_heading(doc, "5.6 Testing Strategy and Methodology", level=2)
    add_paragraph(
        doc,
        "Unit and integration tests are implemented using pytest/unittest style scripts from the tests folder. "
        "Core test focus: preprocessing behavior, training artifact correctness, endpoint behavior, and authentication guards.",
    )
    add_table(
        doc,
        ["Test Case", "Expected Result", "Status"],
        [
            ["GET /health", "Returns status=ok when model is loaded", "Passed"],
            ["POST /predict without auth", "Returns HTTP 401", "Passed"],
            ["POST /predict with valid payload", "Returns prediction and risk_probability", "Passed"],
            ["GET /api/predictions/summary", "Returns count/latest/previous", "Passed"],
            ["POST /api/analyze-report", "Parses report and saves interaction", "Passed"],
            ["Preprocess zero handling", "Converts selected zeros to NaN", "Passed"],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "CHAPTER 6: RESULTS", level=1)
    add_heading(doc, "6.1 Test Reports", level=2)
    model = project_data.get("model", {})
    metrics = model.get("metrics", {}) if isinstance(model, dict) else {}
    dataset_rows = project_data.get("dataset_rows", 0)
    outcome_counts = project_data.get("outcome_counts", {})
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Dataset Rows", str(dataset_rows)],
            ["Dataset Columns", str(project_data.get("dataset_columns", 0))],
            ["Outcome Count (0)", str(outcome_counts.get("0", "N/A"))],
            ["Outcome Count (1)", str(outcome_counts.get("1", "N/A"))],
            ["Selected Model", str(model.get("selected_model", "N/A"))],
            ["Cross-Validation ROC-AUC", f"{model.get('cv_roc_auc', 0):.4f}" if model.get("cv_roc_auc") is not None else "N/A"],
            ["Accuracy", f"{metrics.get('accuracy', 0):.4f}" if metrics else "N/A"],
            ["Precision", f"{metrics.get('precision', 0):.4f}" if metrics else "N/A"],
            ["Recall", f"{metrics.get('recall', 0):.4f}" if metrics else "N/A"],
            ["F1 Score", f"{metrics.get('f1', 0):.4f}" if metrics else "N/A"],
            ["ROC-AUC", f"{metrics.get('roc_auc', 0):.4f}" if metrics else "N/A"],
        ],
    )

    images = project_data.get("images", [])
    if images:
        add_paragraph(doc, "Project Visual Outputs:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for image_path in images:
        doc.add_paragraph()
        doc.add_picture(image_path, width=Inches(5.8))
        caption = f"Figure: {Path(image_path).name}"
        add_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, font_size=11)

    add_heading(doc, "6.2 User Documentation", level=2)
    user_steps = [
        "1. Open the app: http://127.0.0.1:5000/",
        "2. Register a user account or login using existing credentials.",
        "3. Navigate to Predict Risk and fill the clinical input fields.",
        "4. Click Analyze Risk to generate prediction and probability.",
        "5. Use 'Download Latest Predicted File' or 'Download History CSV'.",
        "6. Review model insights and patient record analytics sections.",
    ]
    for step in user_steps:
        add_paragraph(doc, step, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(
        doc,
        "Implemented API Endpoints:",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    route_rows = [[str(index + 1), route] for index, route in enumerate(project_data.get("routes", []))]
    if route_rows:
        add_table(doc, ["Sr", "Route Definition"], route_rows)

    doc.add_page_break()
    add_heading(doc, "CHAPTER 7: CONCLUSION", level=1)
    add_heading(doc, "7.1 Significance of the System", level=2)
    add_paragraph(
        doc,
        "The project demonstrates a complete capstone-ready healthcare analytics workflow: data engineering, model selection, "
        "deployment, user interaction persistence, and downloadable outputs.",
    )
    add_heading(doc, "7.2 Limitations of the System", level=2)
    add_paragraph(
        doc,
        "The current model is trained on a compact benchmark/synthetic-compatible dataset and should be treated as an academic "
        "decision-support prototype, not a clinical diagnostic instrument.",
    )
    add_heading(doc, "7.3 Future Scope of the Project", level=2)
    add_paragraph(
        doc,
        "Future scope includes richer clinical features, model explainability dashboards, role-based access control, cloud deployment, "
        "and integration with secure healthcare information systems.",
    )

    doc.add_page_break()
    add_heading(doc, "REFERENCES / BIBLIOGRAPHY", level=1)
    references = [
        "1. Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
        "2. Cox, D. R. (1958). The Regression Analysis of Binary Sequences.",
        "3. Fawcett, T. (2006). An introduction to ROC analysis. Pattern Recognition Letters.",
        "4. Pedregosa et al. (2011). Scikit-learn: Machine Learning in Python.",
        "5. Flask Documentation. https://flask.palletsprojects.com/",
        "6. SQLite Documentation. https://sqlite.org/docs.html",
        "7. Chart.js Documentation. https://www.chartjs.org/docs/",
        "8. NIDDK / Pima Indians Diabetes Database resources.",
        "9. Hastie, Tibshirani, Friedman. The Elements of Statistical Learning.",
        "10. Project source files: src/app.py, src/train.py, src/utils/*.py, tests/*.py.",
    ]
    for ref in references:
        add_paragraph(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()
    add_heading(doc, "ANNEXURE", level=1)
    add_heading(doc, "1. List of Tables", level=2)
    table_list = [
        "Table A: Requirements Specification",
        "Table B: Planning and Scheduling",
        "Table C: Technology Stack",
        "Table D: Database Design",
        "Table E: Challenges and Solutions",
        "Table F: Test Cases",
        "Table G: Model and Dataset Metrics",
        "Table H: API Route Inventory",
    ]
    for item in table_list:
        add_paragraph(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "2. List of Diagrams/Figures", level=2)
    if images:
        for index, image_path in enumerate(images, start=1):
            add_paragraph(doc, f"Figure {index}: {Path(image_path).name}", align=WD_ALIGN_PARAGRAPH.LEFT)
    else:
        add_paragraph(doc, "No generated figures found in docs/images.", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "3. List of Screenshots", level=2)
    screenshots = [
        "Dashboard Home",
        "Prediction Form and Result Panel",
        "Model Insights Section",
        "Patient Records Table",
        "Prediction History and Download Actions",
    ]
    for shot in screenshots:
        add_paragraph(doc, f"- {shot}", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "4. Plagiarism Report", level=2)
    add_paragraph(
        doc,
        "Attach the official plagiarism report generated as per department submission process.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )


def main() -> None:
    args = parse_args()
    root = Path(__file__).resolve().parents[2]
    project_data = load_project_data(root)

    doc = Document()
    set_document_style(doc)
    set_page_margins(doc.sections[0])

    doc.sections[0].header.paragraphs[0].clear()
    doc.sections[0].footer.paragraphs[0].clear()
    add_cover_pages(doc, args)
    add_index_page(doc)

    main_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_page_margins(main_section)
    set_page_number_start(main_section, 1)
    configure_main_section_header_footer(main_section)

    add_main_content(doc, project_data)
    add_expanded_blackbook_content(doc, root, project_data)

    output_path = root / args.output
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    print(f"Black book generated: {output_path}")
    print("Data used: dataset, model artifact, API routes, full code listings, and project operation manual.")


if __name__ == "__main__":
    main()
